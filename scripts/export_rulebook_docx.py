#!/usr/bin/env python3
"""Export the full DS rulebook (40 rules) + their details to a .docx report."""
import json
import os
import sys
from datetime import datetime

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from agents.loader import load_rulebook
from dataclasses import asdict

OUT = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "..",
    "docs",
    "Rulebook_DS-01_to_DS-40.docx",
)


def main() -> None:
    rules = load_rulebook()
    assert len(rules) == 40, len(rules)

    doc = Document()

    # -- Base style --
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)

    monospace = doc.styles["Normal"]
    code_style = doc.styles["No Spacing"]

    # -- Title page-ish header --
    title = doc.add_heading("Rulebook Drift Monitor — DS Rulebook", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    meta = doc.add_paragraph()
    meta_run = meta.add_run(
        f"Drift Sentinel rules DS-01..DS-40 · exported {datetime.now().strftime('%d %b %Y %H:%M')}"
    )
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    disclaim = doc.add_paragraph()
    d1 = disclaim.add_run(
        "Synthetic-data disclaimer: "
    )
    d1.bold = True
    d2 = disclaim.add_run(
        "The DS rulebook is a paraphrased demonstration schema based on FATF (2020) red-flag "
        "indicators and FIC Directive 9 / FIC Act obligations; rule text is not verbatim from source "
        "and must be verified against the source documents before any production use."
    )
    d2.font.size = Pt(8.5)
    d2.font.color.rgb = RGBColor(0x8A, 0x8A, 0x8A)

    # -- Overview table --
    doc.add_heading("Overview", level=1)
    row_labels = [
        ("Rules", str(len(rules))),
        ("Rule types", "indicator 33 · obligation 4 · mixed 3"),
    ]
    cats: dict[str, int] = {}
    sevs: dict[str, int] = {}
    for r in rules:
        cats[r.category] = cats.get(r.category, 0) + 1
        sevs[r.risk_severity] = sevs.get(r.risk_severity, 0) + 1
    row_labels.append(("Categories", ", ".join(f"{k} {v}" for k, v in sorted(cats.items()))))
    row_labels.append(("Risk severity", ", ".join(f"{k} {v}" for k, v in sorted(sevs.items()))))

    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in row_labels:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    for cell in table.columns[0].cells:
        for p in cell.paragraphs:
            p.runs[0].bold = True

    for r in rules:
        add_rule(doc, r)

    doc.save(OUT)
    print(f"Wrote {OUT} ({os.path.getsize(OUT)} bytes, {len(rules)} rules)")


def add_rule(doc: Document, rule) -> None:
    d = asdict(rule)

    doc.add_heading(f"{d['id']} · {d['name']}", level=1)

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = True

    def row(label: str, value, monospace: bool = False) -> None:
        cells = table.add_row().cells
        cells[0].text = label
        cells[0].width = Inches(1.9)
        lp = cells[0].paragraphs[0]
        lp.runs[0].bold = True
        lp.runs[0].font.size = Pt(9.5)
        text = value if isinstance(value, str) else json.dumps(value, indent=2, ensure_ascii=False)
        text = text or "—"
        p = cells[1].paragraphs[0]
        run = p.add_run(text)
        run.font.size = Pt(9.5)
        if monospace:
            _mono(run)

    row("Category", d["category"])
    row("Rule type", d["rule_type"])
    row("Risk severity", d["risk_severity"])
    row("AI-relevant", "Yes" if d["ai_relevant"] else "No")
    row("Source", d["source"] or "—")
    row("Source reference", d["source_ref"] or "—")
    row("Trigger (natural language)", d["trigger"] or "—")
    row("Full rule text", d["text"])
    row("Keywords", ", ".join(d["keywords"]) if d["keywords"] else "—")
    row("MITRE ATLAS mapping", ", ".join(d["mitre_atlas"]) if d["mitre_atlas"] else "—")
    row(
        "Capability primitives",
        "; ".join(d["capability_primitives"]) if d["capability_primitives"] else "—",
    )
    if d["signal_signature"]:
        row("Signal signature", json.dumps(d["signal_signature"], indent=2, ensure_ascii=False), monospace=True)
    if d["source_finding"]:
        row("Source finding", d["source_finding"])
    if d["institutionalised"]:
        row("Institutionalised", "Yes")
    row("Trigger schema (structured)", d["trigger_schema"] or {}, monospace=True)


def _mono(run) -> None:
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)


if __name__ == "__main__":
    main()