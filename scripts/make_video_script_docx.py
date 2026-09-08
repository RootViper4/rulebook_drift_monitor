#!/usr/bin/env python3
"""Generate the hackathon demo-video script (.docx) for the Rulebook Drift Monitor."""
import os
import sys

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

OUT = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "..",
    "docs",
    "Rulebook_Drift_Monitor_Video_Script.docx",
)

BLUE = RGBColor(0x43, 0x61, 0xEE)
DARK = RGBColor(0x1A, 0x1A, 0x2E)
GREY = RGBColor(0x6B, 0x72, 0x80)
GREEN = RGBColor(0x1B, 0x74, 0x37)

# (cue, scene / visual, narration, on-screen text)
STORYBOARD = [
    (
        "0:00 – 0:15 · Hook",
        "Title card, then a fast cut of a live scan running in the app — log lines streaming.",
        "Every new AI-powered scam makes last quarter's AML rulebook a little more obsolete. "
        "Attacks are written by machines that adapt in days. Rulebooks are updated by humans who move in months. "
        "That gap is rulebook drift.",
        "“Every new AI scam makes last quarter's rulebook a little more obsolete.”  ·  Rulebook Drift Monitor",
    ),
    (
        "0:15 – 0:40 · The problem",
        "Static PDF of a rule on screen, then a 'fraud events' timeline accelerating vs a slow 'rule update' marker.",
        "Regulators ask institutions: show us your AML controls still work. Institutions answer with a static rulebook — "
        "a fixed list of indicators that AI-enabled fraud is re-engineered to sidestep. Nobody can prove controls are still "
        "effective against threats they haven't even seen yet. So we built the Rulebook Drift Monitor: a continuous "
        "stress-test for an institution's AML rulebook, against both the attacks that have happened and the ones AI is inventing next.",
        "Problem: static rulebooks vs AI-paceset fraud.\nSolution: continuously stress-test the rulebook against known AND emerging AI threats.",
    ),
    (
        "0:40 – 1:05 · Known Attacks (reconciliation arm)",
        "App open on Review tab → 'Known Attacks'. Pointer hovers a finding; its evaded rules highlight.",
        "The first arm attacks the rulebook with what has already happened. Twelve documented AI capability primitives — "
        "deepfake video-KYC bypass, voice-clone vishing, synthetic identity documents — are each framed as test cases and stepped "
        "through all forty rules by a deterministic engine. Every rule that stays silent while a real attack passes is a provable gap.",
        "Known Attacks: 12 documented AI attacks × 40 rules · every silent rule = a verifiable gap.",
    ),
    (
        "1:05 – 1:30 · Emerging Threats (generation arm)",
        "Switch to 'Emerging Threats' tab; watch novel scenarios appear with source badges (🤖 AI-written).",
        "The second arm does what a PDF can't — it spawns threats nobody has reported yet. Our generation agent combines the "
        "capability primitives into novel evasion scenarios, the model writes each one, and the same engine tests it for real. "
        "Each finding is honestly labelled: genuinely AI-written, a fallback scenario, or just a test case.",
        "Emerging Threats: the system generates novel AI evasion scenarios itself and tests them for real.",
    ),
    (
        "1:30 – 1:55 · Guardrails — critic, red team, human gate",
        "Clip: critic discards a finding (VERDICT: thrown out). Then click 'Try to fool it' probe. Then a finding waiting under 'Awaiting human approval'.",
        "Nothing gets promoted blindly. An independent critic re-tests every claimed gap and discards anything it can't reproduce. "
        "A live red-team probe tries to fool the checker itself — just to prove it can't be gamed. And no AI-invented rule is ever "
        "instituted automatically: it lands in front of a named analyst, who accepts, amends, or rejects it — with a reason on record.",
        "Guardrails: critic verifies · red-team probe · human approval gate — AI proposes, humans decide.",
    ),
    (
        "1:55 – 2:10 · Standing over the whole picture",
        "Dashboard: Drift Index / coverage cards; then Forecast chart animating six review cycles.",
        "Above it all, the Drift Index — zero means full coverage, a hundred means maximum drift exposure — and a forecast that "
        "projects where the rulebook drifts next. The forecast is honest about what it is: a decision-support projection, not a prediction. "
        "Every run is an evidence trail, not an opinion.",
        "Drift Index 0–100 = how far the rulebook has drifted from full coverage. Forecast = decision support, clearly labelled as such.",
    ),
    (
        "2:10 – 2:30 · Audit trail & accountability",
        "Audit page scrolling: run ids, timestamps, decisions, analyst reasons.",
        "Every decision is logged: the run id, node, message, rulebook version, what each analyst approved and why. "
        "So when a supervisor asks 'who changed this rule and based on what evidence?' — the answer is on screen, in full.",
        "Rules → attacks → test → evidence → finding → human decision → audit trail.",
    ),
    (
        "2:30 – 2:50 · Why this is a regulatory tool, not an AI toy",
        "Closing card: 'Rulebook Drift Monitor' + one line mission and logos/text FIC / FSCA.",
        "We're not claiming our AI decides whether an institution is compliant — that's the regulator's job. "
        "We're giving the regulator and the compliance team the thing they're missing: proof that the rulebook is still tested, "
        "traceable, and accountable against AI-paceset threats. Continuously stress-test the AML control framework. "
        "Trace every test back to a regulatory obligation. Require human validation. Keep the audit trail.",
        "Mission: stress-test the AML/CFT rulebook against known + emerging AI threats · trace to obligations · human-validated · auditable.",
    ),
]

# Full narration block for straight-through recording.
NARRATION_ONLY = [
    "Every new AI-powered scam makes last quarter's AML rulebook a little more obsolete. Attacks are written by machines that "
    "adapt in days; rulebooks are updated by humans who move in months. That gap is rulebook drift.",

    "Regulators ask institutions to show their AML controls still work. Institutions answer with a static rulebook — a fixed list "
    "of indicators that AI-enabled fraud is re-engineered to sidestep. Nobody can prove controls are still effective against "
    "threats they haven't even seen yet. So we built the Rulebook Drift Monitor: a continuous stress-test for an institution's "
    "AML rulebook against the attacks that have happened and the ones AI is inventing next.",

    "First arm — Known Attacks. Twelve documented AI capability primitives, from deepfake video-KYC bypass to voice-clone vishing, "
    "are each framed as test cases and stepped through all forty rules by a deterministic engine. Every rule that stays silent "
    "while a real attack passes is a provable gap.",

    "Second arm — Emerging Threats. The system spawns threats nobody has reported yet: it combines the capability primitives into "
    "novel evasion scenarios, writes each one, and tests it for real. And every finding is honestly labelled — genuinely "
    "AI-written, a fallback, or just a test case.",

    "Nothing gets promoted blindly. An independent critic re-tests every claimed gap and discards anything it can't reproduce. A "
    "live red-team probe tries to fool the checker itself. And no AI-invented rule is ever instituted automatically — it lands in "
    "front of a named analyst who accepts, amends, or rejects it, with a reason on record.",

    "Above it all, the Drift Index measures how far the rulebook has drifted from full coverage, and a forecast projects where "
    "it drifts next — clearly labelled decision support, not a prediction.",

    "Every decision is logged: run id, node, message, rulebook version, what each analyst approved and why. When a supervisor "
    "asks who changed this rule, and based on what evidence, the answer is on screen, in full.",

    "We're not claiming AI decides whether an institution is compliant — that's the regulator's job. We give supervisors and "
    "compliance teams the thing they're missing: proof the rulebook is still tested, traceable, and accountable against "
    "AI-paceset threats. Stress-test the framework. Trace to obligations. Validate with humans. Keep the audit trail.",
]


def _style_para(p, size=11, bold=False, color=None, italic=False):
    for run in p.runs:
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        if color is not None:
            run.font.color.rgb = color


def main() -> None:
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # Title
    t = doc.add_paragraph()
    t.add_run("Rulebook Drift Monitor")
    _style_para(t, size=26, bold=True, color=DARK)
    sub = doc.add_paragraph("Demo-video script · CDIR Global “Agentic Regulator” Hackathon")
    _style_para(sub, size=12, color=GREY)
    meta = doc.add_paragraph(
        "Target length: ~2:50 (max 3:00). Read narration at a measured pace; total narration is ~460 words. "
        "Pair each scene with live screen recordings of the app. Judgment 17 Sep."
    )
    _style_para(meta, size=10, italic=True, color=GREY)

    # The one-sentence arc
    p = doc.add_paragraph()
    p.add_run("The arc in one sentence: ")
    _style_para(p, size=11, bold=True, color=DARK)
    p2 = doc.add_paragraph()
    p2.add_run(
        "AI-paceset fraud is widening the gap between the attack surface and a static AML rulebook; our system continuously "
        "stress-tests that rulebook against both known and emerging AI threats, verifies every gap, keeps humans accountable, "
        "and leaves a complete audit trail."
    )
    _style_para(p2, size=11, italic=True)

    # Storyboard table
    doc.add_heading("Storyboard (scene by scene)", level=1)
    for run in doc.paragraphs[-1].runs:
        run.font.color.rgb = BLUE
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    widths = (Inches(1.0), Inches(1.7), Inches(3.1), Inches(1.6))
    hdr = table.rows[0].cells
    for cell, txt in zip(hdr, ("Cue", "Screen / visual", "Narration (voice-over)", "On-screen text")):
        cell.text = ""
        rr = cell.paragraphs[0].add_run(txt)
        rr.bold = True
        rr.font.size = Pt(10)
        rr.font.color.rgb = BLUE

    for cue, visual, narration, ost in STORYBOARD:
        cells = table.add_row().cells
        cells[0].text = cue
        cells[1].text = visual
        cells[2].text = narration
        cells[3].text = ost
        for c in cells:
            for para in c.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(9)
            c.width = widths[0]
        cells[0].width = widths[0]
        cells[1].width = widths[1]
        cells[2].width = widths[2]
        cells[3].width = widths[3]
    # set column widths via cell spans
    for ci, w in enumerate(widths):
        for row in table.rows:
            row.cells[ci].width = w

    # Narration-only block
    doc.add_heading("Narration only (record straight through)", level=1)
    for run in doc.paragraphs[-1].runs:
        run.font.color.rgb = BLUE
    for para in NARRATION_ONLY:
        p = doc.add_paragraph(para)
        p.paragraph_format.space_after = Pt(8)
        _style_para(p, size=11)

    # Recording tips
    doc.add_heading("Recording & cut notes", level=1)
    for run in doc.paragraphs[-1].runs:
        run.font.color.rgb = BLUE
    tips = [
        "Record narration first, then cut visuals to the audio. OBS at 1080p, browser zoomed so text is legible.",
        "Original scan footage is best: trigger a run in the Review tab, hover the cursor over the two tabs (Known Attacks / Emerging Threats), the Drift Index cards, the red-team 'Try to fool it' probe, a 'thrown out' verdict, and the Audit page.",
        "Point at the honest labels: 'AI-written' (🤖) vs fallback vs 'test case' badges — judges read this as maturity.",
        "Read the Drift Index line carefully: 'zero means full coverage, a hundred means maximum drift exposure'.",
        "Close with the mission sentence, unbroken, over the product card.",
        "Add a subtitle burn for the on-screen-text column at the listed cues.",
        "Keep the demo ≤3:00 total; this trims naturally by shortening scene 2 if it runs long.",
    ]
    for tip in tips:
        p = doc.add_paragraph(tip, style="List Bullet")
        _style_para(p, size=11)

    # Summary of key numbers (verifiable against the code/data)
    doc.add_heading("Numbers to keep consistent (source of truth)", level=1)
    for run in doc.paragraphs[-1].runs:
        run.font.color.rgb = BLUE
    facts = [
        "40 rules (DS-01..DS-40) — rulebook derived from FATF (2020) virtual-asset red-flag indicators and FIC Directive 9 / FIC Act obligations.",
        "Known Attacks arm: 12 AI capability primitives (CP-01..CP-12), stepped through every rule by the deterministic rule engine.",
        "Emerging Threats arm: the generation agent composes up to 5 novel scenarios per run; each is tagged llm / deterministic_fallback / probe.",
        "Guardrails: critic re-verification with discard; live red-team probe ('Try to fool it'); named-analyst human approval gate; no auto-institution of rules.",
        "Drift Index 0–100 (0 full coverage → 100 max drift exposure); forecast = decision-support projection, labelled as such.",
        "Audit trail: run id, timestamp, node, message, rulebook version, analyst decision + reason.",
        "Positioning: control-testing / rulebook stress-testing prototype — explicitly NOT an AI compliance verdict.",
    ]
    for f in facts:
        p = doc.add_paragraph(f, style="List Bullet")
        _style_para(p, size=10.5)

    doc.save(OUT)
    print(f"Wrote {OUT} ({os.path.getsize(OUT)} bytes)")


if __name__ == "__main__":
    main()